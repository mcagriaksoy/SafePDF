"""
PDF OCR conversion module for SafePDF.
Handles scanned/image-based PDF text extraction using Tesseract OCR.
"""

from pathlib import Path
from typing import Tuple

import pypdfium2 as pdfium

from SafePDF.logger.logging_config import get_logger

try:
    import pytesseract
    from pytesseract import TesseractNotFoundError
except ImportError:
    pytesseract = None
    TesseractNotFoundError = RuntimeError


class PDFOCRConverter:
    """Class handling OCR-based PDF conversions."""

    OCR_LANG_MAP = {
        "en": "eng",
        "de": "deu+eng",
        "tr": "tur+eng",
    }

    def __init__(
        self,
        progress_callback=None,
        status_callback=None,
        language_manager=None,
        atomic_write_file=None,
        atomic_write_via_path=None,
    ):
        self.progress_callback = progress_callback
        self.status_callback = status_callback
        self.language_manager = language_manager
        self._atomic_write_file = atomic_write_file
        self._atomic_write_via_path = atomic_write_via_path
        self._cancel_requested = False
        self.logger = get_logger("SafePDF.PDFOCR")
        self._tesseract_checked = False

    def update_progress(self, value):
        """Update progress if callback is available."""
        if self.progress_callback:
            self.progress_callback(value)

    def update_status(self, message):
        """Update live status if callback is available."""
        if self.status_callback:
            self.status_callback(message)

    def _lang_text(self, key: str, default: str) -> str:
        """Return a localized string when a language manager is available."""
        if self.language_manager:
            return self.language_manager.get(key, default)
        return default

    def request_cancel(self):
        """Request cancellation of a running operation."""
        self._cancel_requested = True

    def _get_ocr_langs(self):
        if not self.language_manager:
            return "eng"
        return self.OCR_LANG_MAP.get(getattr(self.language_manager, "lang", "en"), "eng")

    def _configure_tesseract_binary(self):
        """Use the default Windows Tesseract install if it exists."""
        if pytesseract is None or pytesseract.pytesseract.tesseract_cmd != "tesseract":
            return

        candidate_paths = [
            Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
            Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
        ]
        for candidate in candidate_paths:
            if candidate.exists():
                pytesseract.pytesseract.tesseract_cmd = str(candidate)
                return

    def _ensure_ocr_available(self):
        """Validate that pytesseract is importable and the Tesseract binary is installed."""
        if pytesseract is None:
            raise ImportError("pytesseract is not installed")

        self._configure_tesseract_binary()

        if self._tesseract_checked:
            return

        self.update_status(self._lang_text("ocr_status_loading_model", "Loading OCR model..."))
        pytesseract.get_tesseract_version()
        self._tesseract_checked = True
        self.update_status(self._lang_text("ocr_status_model_ready", "OCR model ready."))

    def _extract_page_text_chunks(self, input_path: str, dpi: int = 300):
        """Run OCR on each PDF page and return per-page text chunks."""
        pdf = pdfium.PdfDocument(input_path)
        try:
            total_pages = len(pdf)
            ocr_langs = self._get_ocr_langs()
            self._ensure_ocr_available()
            scale = max(dpi, 72) / 72.0
            text_chunks = []

            for page_num in range(total_pages):
                page_index = page_num + 1
                self.update_status(
                    self._lang_text(
                        "ocr_status_rendering_page",
                        "Rendering page {page} of {total}...",
                    ).format(page=page_index, total=total_pages)
                )
                self.update_progress(int((page_num / total_pages) * 100))
                if self._cancel_requested:
                    return False, self._lang_text("op_word_cancelled", "Operation cancelled")

                page = pdf[page_num]
                pil_image = page.render(scale=scale).to_pil()
                self.update_status(
                    self._lang_text(
                        "ocr_status_recognizing_page",
                        "Recognizing text on page {page} of {total}...",
                    ).format(page=page_index, total=total_pages)
                )
                text = pytesseract.image_to_string(pil_image, lang=ocr_langs).strip()
                text_chunks.append((page_index, text))
                self.update_progress(int((page_index / total_pages) * 100))

            return True, text_chunks
        finally:
            pdf.close()

    def _save_text_output(self, text_chunks, output_path: str):
        """Write OCR text output to a TXT file."""
        self.update_status(self._lang_text("ocr_status_saving_txt", "Saving OCR text as TXT..."))
        text_content = (
            "\n".join(f"--- Page {page_num} ---\n{text.strip()}\n" for page_num, text in text_chunks).strip() + "\n"
        )

        def _write_text(tmpf):
            tmpf.write(text_content.encode("utf-8"))

        if self._atomic_write_file:
            self._atomic_write_file(output_path, _write_text)
        else:
            with open(output_path, "wb") as out_file:
                _write_text(out_file)

    def _save_docx_output(self, text_chunks, output_path: str):
        """Write OCR text output to a DOCX file."""
        self.update_status(self._lang_text("ocr_status_saving_docx", "Saving OCR text as DOCX..."))
        try:
            from docx import Document
        except ImportError:
            return False, self._lang_text(
                "op_docx_unavailable",
                "python-docx not installed. Please install with: pip install python-docx",
            )

        doc = Document()
        doc.add_heading("OCR Output", level=1)
        for page_num, text in text_chunks:
            doc.add_heading(f"Page {page_num}", level=2)
            doc.add_paragraph(text if text.strip() else "[No text detected on this page]")

        def _save_docx(tmp_path):
            doc.save(tmp_path)

        if self._atomic_write_via_path:
            self._atomic_write_via_path(output_path, _save_docx)
        else:
            doc.save(output_path)
        return True, None

    def pdf_ocr_to_txt(self, input_path: str, output_path: str, dpi: int = 300) -> Tuple[bool, str]:
        """Run OCR on each PDF page and save the extracted text to a TXT file."""
        if not pdfium:
            return False, self._lang_text("op_pdfium_unavailable", "pypdfium2 not available")

        if pytesseract is None:
            return False, self._lang_text(
                "op_ocr_unavailable",
                "OCR dependencies missing. Install with: pip install SafePDF[ocr] and install Tesseract OCR.",
            )

        try:
            success, result = self._extract_page_text_chunks(input_path, dpi=dpi)
            if not success:
                return False, result
            self._save_text_output(result, output_path)

            success_msg = self._lang_text("op_ocr_success", "OCR text extracted to {output_path}")
            return True, success_msg.format(output_path=output_path)

        except TesseractNotFoundError:
            return False, self._lang_text("op_ocr_engine_missing", "OCR engine not found.")
        except Exception as e:
            self.logger.error("OCR conversion failed", exc_info=True)
            error_msg = self._lang_text("op_ocr_failed", "PDF OCR failed: {error}")
            return False, error_msg.format(error=str(e))

    def pdf_ocr_to_docx(self, input_path: str, output_path: str, dpi: int = 300) -> Tuple[bool, str]:
        """Run OCR on each PDF page and save the extracted text to a DOCX file."""
        if not pdfium:
            return False, self._lang_text("op_pdfium_unavailable", "pypdfium2 not available")

        if pytesseract is None:
            return False, self._lang_text(
                "op_ocr_unavailable",
                "OCR dependencies missing. Install with: pip install SafePDF[ocr] and install Tesseract OCR.",
            )

        try:
            success, result = self._extract_page_text_chunks(input_path, dpi=dpi)
            if not success:
                return False, result
            success, error = self._save_docx_output(result, output_path)
            if not success:
                return False, error

            success_msg = self._lang_text("op_word_success", "PDF converted to DOCX document: {output_path}")
            return True, success_msg.format(output_path=output_path)
        except TesseractNotFoundError:
            return False, self._lang_text("op_ocr_engine_missing", "OCR engine not found.")
        except Exception as e:
            self.logger.error("OCR DOCX conversion failed", exc_info=True)
            error_msg = self._lang_text("op_ocr_failed", "PDF OCR failed: {error}")
            return False, error_msg.format(error=str(e))
